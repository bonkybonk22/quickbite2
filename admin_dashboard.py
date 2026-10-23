import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font

from database import (
    VERIFICATION_ACCEPTED, VERIFICATION_DECLINED, accounts, db,
)
from ui_theme import (
    BG, CARD_BG, GOLD, NAVY, ORANGE, RED_ACCENT, TEXT_DARK, TEXT_MUTED,
)


class AdminDashboardMixin:
    def build_ingredients_content(self, parent):
        tk.Label(parent, text="INGREDIENT STOCK MANAGEMENT", bg=BG, fg=TEXT_DARK,
                 font=("Segoe UI", 14, "bold")).pack(anchor="w", pady=(0, 12))

        ing_card = tk.Frame(parent, bg=CARD_BG, highlightbackground="#e2e8f0", highlightthickness=1)
        ing_card.pack(fill="both", expand=True, padx=0, pady=0)

        top_act = tk.Frame(ing_card, bg=CARD_BG)
        top_act.pack(fill="x", padx=20, pady=18)
        tk.Label(
            top_act, text="Monitor ingredient availability and stock levels.",
            bg=CARD_BG, fg=TEXT_MUTED, font=("Segoe UI", 9),
        ).pack(side="left")

        tk.Button(top_act, text="+ Restock Ingredient", bg=RED_ACCENT, fg="white",
                  font=("Segoe UI", 9, "bold"), relief="flat", padx=14, pady=8,
                  cursor="hand2", command=self.restock).pack(side="right")

        self.ing_tree = ttk.Treeview(
            ing_card, columns=("id", "name", "status", "stock"),
            show="headings", height=12,
        )
        
        col_widths = (70, 230, 160, 110)
        col_names = ("ID", "INGREDIENT", "STATUS", "STOCK LEVEL")
        for c, w, name in zip(("id", "name", "status", "stock"), col_widths, col_names):
            self.ing_tree.heading(c, text=name)
            self.ing_tree.column(
                c, width=w, minwidth=60, stretch=c == "name",
                anchor="e" if c == "stock" else "w",
            )
        
        self.ing_tree.pack(fill="both", expand=True, padx=20, pady=(0, 18))

    def build_settings_content(self, parent):
        tk.Label(parent, text="SYSTEM SETTINGS", bg=BG, fg=TEXT_DARK,
                 font=("Segoe UI", 14, "bold")).pack(anchor="w", pady=(0, 12))

        settings_card = tk.Frame(parent, bg=CARD_BG, highlightbackground="#e2e8f0", highlightthickness=1)
        settings_card.pack(fill="x", padx=0, pady=(0, 16))

        inner = tk.Frame(settings_card, bg=CARD_BG)
        inner.pack(fill="x", padx=24, pady=22)
        inner.columnconfigure(1, weight=1)

        tk.Label(
            inner, text="BUSINESS DETAILS", bg=CARD_BG, fg=RED_ACCENT,
            font=("Segoe UI", 10, "bold"),
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))

        tk.Label(inner, text="Business Name", bg=CARD_BG, fg=TEXT_DARK,
                 font=("Segoe UI", 9, "bold")).grid(row=1, column=0, sticky="w", pady=6)
        self.biz_name_entry = ttk.Entry(inner, font=("Segoe UI", 10))
        self.biz_name_entry.grid(row=1, column=1, sticky="ew", padx=(20, 0), pady=6)
        self.biz_name_entry.insert(0, db.settings["business_name"])

        tk.Label(inner, text="Receipt Header Subtitle", bg=CARD_BG, fg=TEXT_DARK,
                 font=("Segoe UI", 9, "bold")).grid(row=2, column=0, sticky="w", pady=6)
        self.biz_sub_entry = ttk.Entry(inner, font=("Segoe UI", 10))
        self.biz_sub_entry.grid(row=2, column=1, sticky="ew", padx=(20, 0), pady=6)
        self.biz_sub_entry.insert(0, db.settings["receipt_header"])

        tk.Label(
            inner, text="INVENTORY CONTROLS", bg=CARD_BG, fg=RED_ACCENT,
            font=("Segoe UI", 10, "bold"),
        ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(18, 8))

        tk.Label(inner, text="Low Stock Threshold", bg=CARD_BG, fg=TEXT_DARK,
                 font=("Segoe UI", 9, "bold")).grid(row=4, column=0, sticky="w", pady=6)
        self.low_thresh_entry = ttk.Entry(inner, font=("Segoe UI", 10))
        self.low_thresh_entry.grid(row=4, column=1, sticky="ew", padx=(20, 0), pady=6)
        self.low_thresh_entry.insert(0, str(db.settings["low_threshold"]))

        tk.Label(inner, text="High Stock Threshold", bg=CARD_BG, fg=TEXT_DARK,
                 font=("Segoe UI", 9, "bold")).grid(row=5, column=0, sticky="w", pady=6)
        self.high_thresh_entry = ttk.Entry(inner, font=("Segoe UI", 10))
        self.high_thresh_entry.grid(row=5, column=1, sticky="ew", padx=(20, 0), pady=6)
        self.high_thresh_entry.insert(0, str(db.settings["high_threshold"]))

        tk.Button(
            inner, text="Save Settings", bg=RED_ACCENT, fg="white",
            font=("Segoe UI", 9, "bold"), relief="flat", padx=18, pady=8,
            cursor="hand2", command=self.save_settings,
        ).grid(row=6, column=1, sticky="e", padx=(20, 0), pady=(14, 0))

        # Sales Summary & History Management Card
        stats_card = tk.Frame(parent, bg=CARD_BG, highlightbackground="#e2e8f0", highlightthickness=1)
        stats_card.pack(fill="x", padx=0, pady=0)

        stats_inner = tk.Frame(stats_card, bg=CARD_BG)
        stats_inner.pack(fill="x", padx=20, pady=20)

        tk.Label(stats_inner, text="SALES SUMMARY & HISTORY", bg=CARD_BG, fg=TEXT_DARK, font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(0, 10))
        
        self.sales_stat_lbl = tk.Label(stats_inner, text="", bg=CARD_BG, fg=TEXT_MUTED, font=("Segoe UI", 10))
        self.sales_stat_lbl.pack(anchor="w", pady=4)

        tk.Button(stats_inner, text="🗑 Clear Order History", bg="#d62828", fg="white", font=("Segoe UI", 9, "bold"),
                  relief="flat", padx=14, pady=6, cursor="hand2", command=self.clear_history).pack(anchor="w", pady=(12, 0))

    def build_monthly_sales_content(self, parent):
        tk.Label(
            parent, text="MONTHLY FOOD SALES REPORT", bg=BG, fg=TEXT_DARK,
            font=("Segoe UI", 14, "bold"),
        ).pack(anchor="w", pady=(0, 12))

        report_card = tk.Frame(parent, bg=CARD_BG, highlightbackground="#e2e8f0", highlightthickness=1)
        report_card.pack(fill="both", expand=True)

        report_header = tk.Frame(report_card, bg=CARD_BG)
        report_header.pack(fill="x", padx=20, pady=(16, 10))
        tk.Label(
            report_header, text="Sales by food item", bg=CARD_BG,
            fg=TEXT_DARK, font=("Segoe UI", 11, "bold"),
        ).pack(side="left")

        report_controls = tk.Frame(report_card, bg=CARD_BG)
        report_controls.pack(fill="x", padx=20, pady=(0, 10))
        today = datetime.now()
        self.report_month = ttk.Combobox(
            report_controls, state="readonly", width=12,
            values=tuple(datetime(2000, month, 1).strftime("%B") for month in range(1, 13)),
        )
        self.report_month.current(today.month - 1)
        self.report_month.pack(side="left", padx=(0, 8))
        self.report_year = ttk.Entry(report_controls, width=8)
        self.report_year.insert(0, str(today.year))
        self.report_year.pack(side="left", padx=(0, 8))
        tk.Button(
            report_controls, text="Generate Report", bg=ORANGE, fg="white",
            font=("Segoe UI", 9, "bold"), relief="flat", padx=12, pady=6,
            cursor="hand2", command=self.generate_monthly_sales_report,
        ).pack(side="left")
        tk.Button(
            report_controls, text="Export to Excel", bg=NAVY, fg="white",
            font=("Segoe UI", 9, "bold"), relief="flat", padx=12, pady=6,
            cursor="hand2", command=self.export_monthly_sales_report,
        ).pack(side="left", padx=(8, 0))

        report_columns = ("item", "quantity", "revenue")
        self.monthly_sales_tree = ttk.Treeview(
            report_card, columns=report_columns, show="headings", height=5,
        )
        self.monthly_sales_tree.heading("item", text="FOOD ITEM")
        self.monthly_sales_tree.heading("quantity", text="QUANTITY SOLD")
        self.monthly_sales_tree.heading("revenue", text="SALES")
        self.monthly_sales_tree.column("item", width=260, anchor="w")
        self.monthly_sales_tree.column("quantity", width=140, anchor="center")
        self.monthly_sales_tree.column("revenue", width=160, anchor="e")
        self.monthly_sales_tree.pack(fill="both", expand=True, padx=20, pady=(0, 8))
        self.monthly_sales_summary = tk.Label(
            report_card, text="", bg=CARD_BG, fg=TEXT_DARK,
            font=("Segoe UI", 9, "bold"),
        )
        self.monthly_sales_summary.pack(anchor="e", padx=20, pady=(0, 14))
        self.generate_monthly_sales_report()

    def generate_monthly_sales_report(self):
        try:
            year = int(self.report_year.get().strip())
            month = self.report_month.current() + 1
            if not 1 <= year <= 9999 or month == 0:
                raise ValueError
        except ValueError:
            messagebox.showerror(
                "Invalid report period", "Select a month and enter a valid year.", parent=self
            )
            return

        sales = db.get_monthly_sales(year, month)
        self.monthly_sales_tree.delete(*self.monthly_sales_tree.get_children())
        total_quantity = 0
        total_revenue = 0
        for item in sales:
            total_quantity += item["quantity"]
            total_revenue += item["revenue"]
            self.monthly_sales_tree.insert(
                "", "end",
                values=(item["name"], item["quantity"], f"₱{item['revenue']:.2f}"),
            )

        self.monthly_sales_summary.config(
            text=(
                f"{self.report_month.get()} {year}  |  "
                f"Total items sold: {total_quantity}  |  "
                f"Total sales: ₱{total_revenue:.2f}"
            )
        )

    def export_monthly_sales_report(self):
        try:
            year = int(self.report_year.get().strip())
            month = self.report_month.current() + 1
            if not 1 <= year <= 9999 or month == 0:
                raise ValueError
        except ValueError:
            messagebox.showerror(
                "Invalid report period", "Select a month and enter a valid year.", parent=self
            )
            return

        filename = f"monthly_sales_{year}_{month:02d}.xlsx"
        filepath = filedialog.asksaveasfilename(
            parent=self,
            title="Export Monthly Sales Report",
            defaultextension=".xlsx",
            initialfile=filename,
            filetypes=(("Excel Workbook", "*.xlsx"),),
        )
        if not filepath:
            return

        sales = db.get_monthly_sales(year, month)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Monthly Sales"
        worksheet.append(["Monthly Food Sales Report"])
        worksheet["A1"].font = Font(bold=True, size=16)
        worksheet.append([f"{self.report_month.get()} {year}"])
        worksheet.append(["Food Item", "Quantity Sold", "Sales"])
        for cell in worksheet[3]:
            cell.font = Font(bold=True)

        total_quantity = 0
        total_revenue = 0
        for item in sales:
            total_quantity += item["quantity"]
            total_revenue += item["revenue"]
            worksheet.append([item["name"], item["quantity"], item["revenue"]])

        worksheet.append(["Total", total_quantity, total_revenue])
        for cell in worksheet[worksheet.max_row]:
            cell.font = Font(bold=True)
        for row in worksheet.iter_rows(min_row=4, min_col=3, max_col=3):
            row[0].number_format = '"₱"#,##0.00'
        worksheet.column_dimensions["A"].width = 30
        worksheet.column_dimensions["B"].width = 18
        worksheet.column_dimensions["C"].width = 18

        try:
            workbook.save(filepath)
        except OSError as error:
            messagebox.showerror("Export failed", f"Could not save the Excel file:\n{error}", parent=self)
            return
        messagebox.showinfo("Export complete", f"Monthly sales report saved to:\n{filepath}", parent=self)

    def build_verification_content(self, parent):
        tk.Label(
            parent, text="PENDING ACCOUNT VERIFICATION", bg=BG, fg=TEXT_DARK,
            font=("Segoe UI", 14, "bold"),
        ).pack(anchor="w", pady=(0, 12))

        card = tk.Frame(parent, bg=CARD_BG, highlightbackground="#e2e8f0", highlightthickness=1)
        card.pack(fill="both", expand=True)

        toolbar = tk.Frame(card, bg=CARD_BG)
        toolbar.pack(fill="x", padx=16, pady=16)
        tk.Label(
            toolbar, text="Review pending registrations below.",
            bg=CARD_BG, fg=TEXT_MUTED, font=("Segoe UI", 10),
        ).pack(side="left")
        tk.Button(
            toolbar, text="Refresh", bg=GOLD, fg=TEXT_DARK,
            font=("Segoe UI", 9, "bold"), relief="flat", padx=12, pady=6,
            cursor="hand2", command=self.refresh_pending_accounts,
        ).pack(side="right")

        columns = ("username", "role", "verification_status")
        self.pending_tree = ttk.Treeview(card, columns=columns, show="headings", height=12)
        for column, label, width in (
            ("username", "USERNAME", 220),
            ("role", "ROLE", 180),
            ("verification_status", "STATUS", 160),
        ):
            self.pending_tree.heading(column, text=label)
            self.pending_tree.column(column, width=width, anchor="w")
        self.pending_tree.pack(fill="both", expand=True, padx=16, pady=(0, 12))

        actions = tk.Frame(card, bg=CARD_BG)
        actions.pack(fill="x", padx=16, pady=(0, 16))
        tk.Button(
            actions, text="Accept", bg="#2e9e5b", fg="white",
            font=("Segoe UI", 9, "bold"), relief="flat", padx=18, pady=7,
            cursor="hand2",
            command=lambda: self.review_selected_account(VERIFICATION_ACCEPTED),
        ).pack(side="left", padx=(0, 8))
        tk.Button(
            actions, text="Decline", bg="#d62828", fg="white",
            font=("Segoe UI", 9, "bold"), relief="flat", padx=18, pady=7,
            cursor="hand2",
            command=lambda: self.review_selected_account(VERIFICATION_DECLINED),
        ).pack(side="left")

    def refresh_pending_accounts(self):
        if self.current_role != "Canteen staff" or not hasattr(self, "pending_tree"):
            return
        try:
            pending_accounts = accounts.get_pending_accounts()
        except (OSError, ValueError) as error:
            messagebox.showerror("Unable to load accounts", str(error), parent=self)
            return

        self.pending_tree.delete(*self.pending_tree.get_children())
        for account in pending_accounts:
            self.pending_tree.insert(
                "", "end", iid=account["username"],
                values=(
                    account["username"],
                    account["role"],
                    account["verification_status"],
                ),
            )

    def review_selected_account(self, status):
        if self.current_role != "Canteen staff":
            messagebox.showerror(
                "Access denied", "Only canteen staff can verify accounts.", parent=self
            )
            return
        selection = self.pending_tree.selection()
        if not selection:
            messagebox.showerror(
                "No account selected", "Select a pending account to review.", parent=self
            )
            return
        try:
            accounts.update_verification(self.current_user, selection[0], status)
        except (OSError, ValueError) as error:
            messagebox.showerror("Unable to update account", str(error), parent=self)
            return
        self.refresh_pending_accounts()

    def save_settings(self):
        try:
            business_name = self.biz_name_entry.get().strip()
            receipt_header = self.biz_sub_entry.get().strip()
            low_threshold = int(self.low_thresh_entry.get().strip())
            high_threshold = int(self.high_thresh_entry.get().strip())
            if low_threshold < 0 or high_threshold <= low_threshold:
                raise ValueError
        
            db.settings["business_name"] = business_name
            db.settings["receipt_header"] = receipt_header
            db.settings["low_threshold"] = low_threshold
            db.settings["high_threshold"] = high_threshold
            for ingredient in db.get_ingredients():
                if ingredient["stock"] >= high_threshold:
                    ingredient["status"] = "High"
                elif ingredient["stock"] <= low_threshold:
                    ingredient["status"] = "Low Stock"
                else:
                    ingredient["status"] = "Normal"
            for item in db.get_menu():
                if item["stock"] <= 0:
                    item["status"] = "Out of Stock"
                elif item["stock"] <= low_threshold:
                    item["status"] = "Low Stock"
                else:
                    item["status"] = "Available"
            self.header_main_lbl.config(text=db.settings["business_name"])
            self.header_sub_lbl.config(text=db.settings["receipt_header"])
            messagebox.showinfo("Success", "Settings saved successfully!")
            self.refresh()
        except ValueError:
            messagebox.showerror(
                "Invalid settings",
                "Enter whole-number thresholds. Low stock must be zero or more, "
                "and high stock must be greater than low stock.",
                parent=self,
            )

    def clear_history(self):
        if messagebox.askyesno("Confirm", "Are you sure you want to clear all order history and sales stats?"):
            db.clear_orders()
            self.refresh()
            messagebox.showinfo("Cleared", "Order history cleared.")

    def refresh_ingredients_table(self):
        if hasattr(self, 'ing_tree'):
            self.ing_tree.delete(*self.ing_tree.get_children())
            for ing in db.get_ingredients():
                self.ing_tree.insert(
                    "", "end", iid=str(ing["id"]), tags=(ing["status"],),
                    values=(ing["id"], ing["name"], ing["status"], ing["stock"])
                )

    def refresh_settings_stats(self):
        if hasattr(self, 'sales_stat_lbl'):
            total_orders = len(db.orders)
            total_revenue = sum(o["total"] for o in db.orders)
            self.sales_stat_lbl.config(text=f"Total Completed Orders: {total_orders}\nTotal Sales Revenue: ₱{total_revenue:.2f}")
